# Copyright (c) 2021, Frappe Technologies Pvt. Ltd. and contributors
# For license information, please see license.txt

import frappe
from frappe import _
from frappe.model.document import Document
import openpyxl

class OpenBOMImportTool(Document):
	pass

@frappe.whitelist()
def process_open_bom(import_file):
	file_name = frappe.get_site_path() + import_file
	asms = []
	items = []
	wb = openpyxl.load_workbook(file_name)
	ws = wb.active

	line_count = 0
	errors = []
	for row in ws.iter_rows(values_only=True):
		if line_count < 9:
			line_count += 1
			continue
		if "." not in row[0]:
			try:
				op = frappe.get_doc("Operation", row[2])
				asms.append(row[2])
			except:
				errors.append(row[0] + " : Opertation not found " + row[2])
		else:
			try:
				item = frappe.get_doc("Item", row[1])
				items.append(row[1])
			except:
				errors.append(row[0] + " : Item not found " + row[1])
		line_count += 1
	if len(errors) == 0:
		frappe.msgprint(_("BOM File processed successfully without error.. You can now import the BOM."))
	return errors

@frappe.whitelist()
def import_open_bom(import_file, item_code):
	file_name = frappe.get_site_path() + import_file
	bom = frappe.new_doc("BOM")
	wb = openpyxl.load_workbook(file_name)
	ws = wb.active

	line_count = 0
	operation = None
	op_qty = 0
	bom.item = item_code
	bom.with_operations = 1
	bom.transfer_material_against = "Job Card"
	bom.set_rate_of_sub_assembly_item_based_on_bom = 0
	
	for row in ws.iter_rows(values_only=True):
		if line_count < 9:
			line_count += 1
			continue
		if "." not in row[0]:
			operation = row[2]
			op_qty = frappe.utils.flt(row[3])
			if(op_qty == 0):
				op_qty = 1
			op = bom.append('operations', {})
			op.operation = operation
			op.workstation = "Assembly Bay"
			op.time_in_mins = 10
		elif operation is None:
			frappe.throw("Error : Operation not found at line {}".format(line_count))
		else:
			item = bom.append('items', {})
			item.item_code = row[1]
			item.operation = operation
			item.qty = frappe.utils.flt(row[3]) * op_qty
			item.stock_uom = item.uom
	bom.save()
	for item in bom.items:
		item.stock_uom = item.uom
	bom.save()
	#bom.submit()
	frappe.db.commit()
	frappe.msgprint(_("New BOM: {0} for Item: {1} created.")
					.format(bom.name, bom.item))